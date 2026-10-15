"""Unit tests for `src.engines.book_engine.document_extraction`."""

from __future__ import annotations

from pathlib import Path
from unittest.mock import MagicMock, patch

import docx
import pytest
from ebooklib import epub

from src.engines.book_engine.document_extraction import (
    extract_document_metadata,
    extract_document_pages,
)
from src.engines.book_engine.text_extraction import ExtractedPage

pytestmark = pytest.mark.unit


def _write_epub(path: Path, *, title: str | None, author: str | None) -> None:
    book = epub.EpubBook()
    book.set_identifier("test-book-id")
    if title:
        book.set_title(title)
    if author:
        book.add_author(author)
    book.set_language("en")

    chapter_one = epub.EpubHtml(title="Chapter 1", file_name="chap_01.xhtml", lang="en")
    chapter_one.content = "<html><body><p>Order blocks precede displacement.</p></body></html>"
    chapter_two = epub.EpubHtml(title="Chapter 2", file_name="chap_02.xhtml", lang="en")
    chapter_two.content = "<html><body><p>Liquidity sweeps often reverse.</p></body></html>"

    book.add_item(chapter_one)
    book.add_item(chapter_two)
    book.toc = (chapter_one, chapter_two)
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    book.spine = ["nav", chapter_one, chapter_two]

    epub.write_epub(str(path), book)


def _write_docx(path: Path, *, title: str | None, author: str | None) -> None:
    document = docx.Document()
    document.add_paragraph("Smart money leaves footprints in the order book.")
    document.add_paragraph("Every liquidity pool is eventually swept.")
    # python-docx defaults core_properties.title to "" (falsy, falls back
    # cleanly) but core_properties.author to its own "python-docx" (truthy)
    # -- both must be explicitly cleared to exercise the fallback branch.
    document.core_properties.title = title or ""
    document.core_properties.author = author or ""
    document.save(str(path))


def test_extract_document_pages_raises_for_missing_file(tmp_path: Path) -> None:
    with pytest.raises(FileNotFoundError):
        extract_document_pages(tmp_path / "does_not_exist.txt")


def test_extract_document_pages_raises_for_unsupported_extension(tmp_path: Path) -> None:
    path = tmp_path / "notes.md"
    path.write_text("hello")

    with pytest.raises(ValueError, match="Unsupported document extension"):
        extract_document_pages(path)


def test_extract_document_pages_reads_txt_as_a_single_page(tmp_path: Path) -> None:
    path = tmp_path / "notes.txt"
    path.write_text("Risk management is the foundation of longevity.", encoding="utf-8")

    pages = extract_document_pages(path)

    assert len(pages) == 1
    assert "Risk management" in pages[0].text
    assert pages[0].used_ocr is False


def test_extract_document_pages_reads_docx_paragraphs(tmp_path: Path) -> None:
    path = tmp_path / "book.docx"
    _write_docx(path, title="DOCX Book", author="A. Trader")

    pages = extract_document_pages(path)

    assert len(pages) == 1
    assert "Smart money" in pages[0].text
    assert "liquidity pool" in pages[0].text


def test_extract_document_pages_reads_epub_chapters_stripped_of_html(tmp_path: Path) -> None:
    path = tmp_path / "book.epub"
    _write_epub(path, title="EPUB Book", author="A. Trader")

    pages = extract_document_pages(path)

    # get_items_of_type(ITEM_DOCUMENT) also includes the auto-generated
    # EpubNav (table of contents) page alongside the two real chapters.
    assert len(pages) == 3
    assert "Order blocks precede displacement." in pages[0].text
    assert "<p>" not in pages[0].text
    assert "Liquidity sweeps often reverse." in pages[1].text


def test_extract_document_pages_delegates_pdf_to_extract_pages(tmp_path: Path) -> None:
    path = tmp_path / "book.pdf"
    path.write_bytes(b"irrelevant")

    sentinel_pages = [ExtractedPage(page_number=1, text="sentinel", used_ocr=False)]
    with patch(
        "src.engines.book_engine.document_extraction.extract_pages",
        return_value=sentinel_pages,
    ) as mock_extract_pages:
        result = extract_document_pages(path)

    mock_extract_pages.assert_called_once_with(path)
    assert result == sentinel_pages


def test_extract_document_metadata_falls_back_for_txt(tmp_path: Path) -> None:
    path = tmp_path / "notes.txt"
    path.write_text("no metadata here")

    metadata = extract_document_metadata(path, fallback_title="notes")

    assert metadata.title == "notes"
    assert metadata.author == "unknown"


def test_extract_document_metadata_reads_docx_core_properties(tmp_path: Path) -> None:
    path = tmp_path / "book.docx"
    _write_docx(path, title="DOCX Book", author="A. Trader")

    metadata = extract_document_metadata(path, fallback_title="book")

    assert metadata.title == "DOCX Book"
    assert metadata.author == "A. Trader"


def test_extract_document_metadata_falls_back_when_docx_has_no_core_properties(
    tmp_path: Path,
) -> None:
    path = tmp_path / "book.docx"
    _write_docx(path, title=None, author=None)

    metadata = extract_document_metadata(path, fallback_title="book")

    assert metadata.title == "book"
    assert metadata.author == "unknown"


def test_extract_document_metadata_reads_epub_dc_metadata(tmp_path: Path) -> None:
    path = tmp_path / "book.epub"
    _write_epub(path, title="EPUB Book", author="A. Trader")

    metadata = extract_document_metadata(path, fallback_title="book")

    assert metadata.title == "EPUB Book"
    assert metadata.author == "A. Trader"


def test_extract_document_metadata_falls_back_when_epub_has_no_dc_metadata(
    tmp_path: Path,
) -> None:
    path = tmp_path / "book.epub"
    _write_epub(path, title=None, author=None)

    metadata = extract_document_metadata(path, fallback_title="book")

    assert metadata.title == "book"
    assert metadata.author == "unknown"


def test_extract_document_metadata_reads_pdf_metadata(tmp_path: Path) -> None:
    path = tmp_path / "book.pdf"
    path.write_bytes(b"irrelevant")
    fake_reader = MagicMock()
    fake_reader.metadata = MagicMock(title="PDF Book", author="A. Trader")

    with patch(
        "src.engines.book_engine.document_extraction.pypdf.PdfReader",
        return_value=fake_reader,
    ):
        metadata = extract_document_metadata(path, fallback_title="book")

    assert metadata.title == "PDF Book"
    assert metadata.author == "A. Trader"


def test_extract_document_metadata_falls_back_when_pdf_metadata_is_missing(
    tmp_path: Path,
) -> None:
    path = tmp_path / "book.pdf"
    path.write_bytes(b"irrelevant")
    fake_reader = MagicMock()
    fake_reader.metadata = None

    with patch(
        "src.engines.book_engine.document_extraction.pypdf.PdfReader",
        return_value=fake_reader,
    ):
        metadata = extract_document_metadata(path, fallback_title="book")

    assert metadata.title == "book"
    assert metadata.author == "unknown"
